"""The HTTP layer: upload returns at once, the report fills in behind it. Labeling and checks are stubbed."""

import datetime as dt
import io
import json
import threading
import time
from pathlib import Path
from types import SimpleNamespace

import jsonschema
import openpyxl
import pytest
from fastapi.testclient import TestClient

import app as engine
import pipeline
from checks import label, news

SCHEMA = json.loads((Path(__file__).parents[2] / "schema" / "constant.schema.json").read_text())
RATE, UNITS, OIL, HALF = "Model!B2", "Model!B3", "Model!B4", "Model!B5"


def _workbook(model_date=dt.date(2026, 8, 1), rates=0) -> bytes:
    """EBITDA = units * rate * 0.5 - oil price, plus `rates` more market rows that EBITDA adds up."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Model"
    if model_date:
        ws["A1"], ws["B1"] = "Model date", model_date
    ws["A2"], ws["B2"] = "EUR/USD rate", 1.25
    ws["A3"], ws["B3"] = "Units", 100
    ws["A4"], ws["B4"] = "Oil price", 80
    ws["A5"], ws["B5"] = "EBITDA", f"=B3*B2*0.5-B4+SUM(B6:B{6 + rates})"
    for row in range(6, 6 + rates):
        ws[f"A{row}"], ws[f"B{row}"] = f"Lane {row} rate", float(row)
    data = io.BytesIO()
    wb.save(data)
    return data.getvalue()


def _apply_all(found, wb, graph):
    for c in found:
        c["label"] = wb.context(c["cell"])["row_label"] or ""
        market = "rate" in c["label"] or "price" in c["label"]
        c["role"] = "market" if market and not c["embedded"] else "internal"
        c["driver"] = label.slug(c["label"]) if c["role"] == "market" else None
        c["label_source"] = "llm"


@pytest.fixture
def api(tmp_path, monkeypatch):
    """A client on the app with uploads in tmp_path. Checks wait for api.gate, which starts open, and
    give a live value 20% above the constant; a driver in api.broken makes its check raise."""
    gate = threading.Event()
    gate.set()
    broken: set[str] = set()

    def check(c, built_on, as_of=None):
        assert gate.wait(10)
        if c["driver"] in broken:
            raise RuntimeError("scrape fell over")
        c["scratch"] = True  # a check that writes on its constant must not reach the report
        return {"kind": "news", "status": "stale", "live": c["value"] * 1.2, "evidence": []}

    monkeypatch.setattr(label, "apply_all", _apply_all)
    monkeypatch.setattr(news, "check", check)
    monkeypatch.setattr(engine, "UPLOADS", tmp_path / "uploads")
    monkeypatch.setattr(engine, "REPORTS", {})
    client = TestClient(engine.app)

    def upload(data=None, name="model.xlsx", **params):
        response = client.post("/models", files={"file": (name, data or _workbook())}, params=params)
        assert response.status_code == 200, response.text
        return response.json()["model_id"]

    def report(model_id):
        found = client.get(f"/models/{model_id}").json()
        jsonschema.validate(found, SCHEMA)
        return found

    def wait(model_id, until=lambda r: r["status"] != "checking"):
        deadline = time.monotonic() + 10
        while time.monotonic() < deadline:
            found = report(model_id)
            if until(found):
                return found
            time.sleep(0.005)
        raise AssertionError(f"timed out: {found['status']} {found['progress']}")

    yield SimpleNamespace(client=client, gate=gate, broken=broken, upload=upload, report=report, wait=wait)
    # no background thread may outlive the stubs: it would go on to call the real checks
    gate.set()
    for model_id, found in list(engine.REPORTS.items()):
        if found.get("status") == "checking":
            wait(model_id)


def _by_cell(report):
    return {c["cell"]: c for c in report["constants"]}


def test_upload_returns_while_the_checks_are_still_running(api):
    api.gate.clear()
    model_id = api.upload()

    report = api.wait(model_id, lambda r: r["progress"]["total"])
    assert report["status"] == "checking"
    assert report["progress"] == {"checked": 0, "total": 2}
    assert [c["cell"] for c in report["constants"]] == [RATE, UNITS, OIL, HALF]
    assert [c["label"] for c in report["constants"]] == ["EUR/USD rate", "Units", "Oil price", "EBITDA"]
    assert all(c["check"] is None and c["impact"] == [] for c in report["constants"])
    assert report["outputs"] == [{"cell": HALF, "label": "EBITDA", "value": -17.5}]
    assert report["combined"] == []

    api.gate.set()
    report = api.wait(model_id)
    assert report["status"] == "done"
    assert report["progress"] == {"checked": 2, "total": 2}
    assert [c["cell"] for c in report["constants"][:2]] == [OIL, RATE]  # largest impact first
    constants = _by_cell(report)
    assert constants[RATE]["check"]["live"] == pytest.approx(1.5)
    assert constants[RATE]["impact"] == [{"output": HALF, "delta": pytest.approx(12.5)}]
    assert constants[RATE]["sensitivity"]["effects"][0]["direction"] == "rises"
    assert constants[OIL]["impact"] == [{"output": HALF, "delta": pytest.approx(-16)}]
    assert constants[UNITS]["check"] is None and constants[UNITS]["impact"] == []
    assert report["combined"] == [{"output": HALF, "delta": pytest.approx(-3.5)}]
    assert not any("scratch" in c for c in report["constants"])


def test_built_on_is_read_from_the_workbook_unless_given(api):
    assert api.wait(api.upload())["built_on"] == "2026-08-01"
    assert api.wait(api.upload(_workbook(model_date=None)))["built_on"] == dt.date.today().isoformat()
    assert api.wait(api.upload(built_on="2025-01-31"))["built_on"] == "2025-01-31"
    assert api.wait(api.upload(built_on=""))["built_on"] == "2026-08-01"

    response = api.client.post("/models", files={"file": ("model.xlsx", _workbook())}, params={"built_on": "last year"})
    assert response.status_code == 422


@pytest.mark.parametrize("data", [b"not a workbook", _workbook()[:400], b""])
def test_a_file_that_is_not_a_workbook_is_refused(api, data):
    response = api.client.post("/models", files={"file": ("notes.xlsx", data)})
    assert response.status_code == 400
    assert "workbook" in response.json()["detail"]
    assert engine.REPORTS == {}
    assert list(engine.UPLOADS.iterdir()) == []


def test_the_uploaded_file_can_be_downloaded(api):
    data = _workbook()
    model_id = api.upload(data, name="Q3 plan.xlsx")
    response = api.client.get(f"/models/{model_id}/file")
    assert response.status_code == 200
    assert response.content == data
    assert "Q3%20plan.xlsx" in response.headers["content-disposition"]

    assert api.client.get("/models/nope/file").status_code == 404
    api.wait(model_id)
    (engine.UPLOADS / f"{model_id}.xlsx").unlink()
    assert api.client.get(f"/models/{model_id}/file").status_code == 404


def test_unknown_model_is_404(api):
    assert api.client.get("/models/nope").status_code == 404


def test_model_list_has_status(api):
    api.gate.clear()
    model_id = api.upload()
    summaries = {m["model_id"]: m for m in api.client.get("/models").json()}
    assert summaries[model_id]["status"] == "checking"
    assert summaries[model_id]["constants"] == 4

    api.gate.set()
    api.wait(model_id)
    summaries = {m["model_id"]: m for m in api.client.get("/models").json()}
    assert (summaries[model_id]["status"], summaries[model_id]["stale"]) == ("done", 2)


def test_an_edit_recalculates_impact_and_combined(api):
    model_id = api.upload()
    api.wait(model_id)
    url = f"/models/{model_id}/constants"

    edited = api.client.patch(url, json={"cell": RATE, "role": "internal"}).json()
    assert (edited["check"], edited["impact"], edited["driver"]) == (None, [], None)
    assert (edited["label_source"], edited["confirmed"]) == ("user", True)
    assert api.report(model_id)["combined"] == [{"output": HALF, "delta": pytest.approx(-16)}]

    edited = api.client.patch(url, json={"cell": RATE, "role": "market", "label": "Euro rate"}).json()
    assert (edited["driver"], edited["search"]) == ("euro_rate", ["Euro rate"])
    assert edited["check"]["live"] == pytest.approx(1.5)
    assert edited["impact"] == [{"output": HALF, "delta": pytest.approx(12.5)}]
    report = api.report(model_id)
    assert report["combined"] == [{"output": HALF, "delta": pytest.approx(-3.5)}]
    assert _by_cell(report)[RATE] == edited

    # a constant that was never market gets its sensitivity along with its first check
    edited = api.client.patch(url, json={"cell": UNITS, "role": "market"}).json()
    assert edited["sensitivity"]["effects"][0]["output"] == HALF
    assert edited["impact"] == [{"output": HALF, "delta": pytest.approx(12.5)}]

    assert api.client.patch(url, json={"cell": RATE, "role": "guess"}).status_code == 422
    assert api.client.patch(url, json={"cell": "Model!Z9", "role": "market"}).status_code == 404
    assert api.client.patch("/models/nope/constants", json={"cell": RATE}).status_code == 404


def test_an_edit_is_refused_while_the_model_is_being_checked(api):
    api.gate.clear()
    model_id = api.upload()
    response = api.client.patch(f"/models/{model_id}/constants", json={"cell": RATE, "role": "internal"})
    assert response.status_code == 409
    assert "still being checked" in response.json()["detail"]

    api.gate.set()
    api.wait(model_id)
    assert api.client.patch(f"/models/{model_id}/constants", json={"cell": RATE, "role": "internal"}).status_code == 200


def test_an_edit_still_recalculates_once_the_workbook_has_left_memory(api):
    model_id = api.upload()
    api.wait(model_id)
    pipeline._LOADED.clear()
    pipeline._PATHS.clear()  # as after a restart: the app knows where the upload is

    edited = api.client.patch(f"/models/{model_id}/constants", json={"cell": OIL, "role": "lever"}).json()
    assert edited["impact"] == []
    assert api.report(model_id)["combined"] == [{"output": HALF, "delta": pytest.approx(12.5)}]


def test_a_check_that_raises_leaves_its_constant_unchecked(api):
    api.broken.add("oil_price")
    report = api.wait(api.upload())
    assert report["status"] == "done"
    assert report["progress"] == {"checked": 2, "total": 2}
    constants = _by_cell(report)
    assert constants[OIL]["check"] is None and constants[OIL]["impact"] == []
    assert constants[RATE]["impact"] == [{"output": HALF, "delta": pytest.approx(12.5)}]
    assert report["combined"] == [{"output": HALF, "delta": pytest.approx(12.5)}]


def test_a_stage_that_breaks_marks_the_report_failed(api, monkeypatch):
    def broken(*args):
        raise RuntimeError("recalculation fell over")

    monkeypatch.setattr(pipeline.impact, "combined", broken)
    report = api.wait(api.upload())
    assert report["status"] == "failed"
    assert _by_cell(report)[RATE]["check"]["live"] == pytest.approx(1.5)  # what was computed stays


def test_reads_during_a_run_are_always_whole(api, monkeypatch):
    check = news.check

    def slow(c, built_on, as_of=None):
        time.sleep(0.01)
        return check(c, built_on, as_of)

    monkeypatch.setattr(news, "check", slow)
    model_id = api.upload(_workbook(rates=40))
    seen, errors = [], []

    def hammer():
        client = TestClient(engine.app)
        try:
            while True:
                response = client.get(f"/models/{model_id}")
                assert response.status_code == 200, response.text
                seen.append(response.json())
                assert client.get("/models").status_code == 200
                if seen[-1]["status"] != "checking":
                    return
        except Exception as error:
            errors.append(error)

    readers = [threading.Thread(target=hammer) for _ in range(3)]
    for reader in readers:
        reader.start()
    for reader in readers:
        reader.join(20)

    assert not errors
    assert any(0 < r["progress"]["checked"] < r["progress"]["total"] for r in seen)
    for report in seen:
        jsonschema.validate(report, SCHEMA)
        assert len(report["constants"]) == 44
        checked = [c for c in report["constants"] if c["check"]]
        assert len(checked) == report["progress"]["checked"] <= report["progress"]["total"]
        assert all(c["impact"] and "sensitivity" in c for c in checked)  # a check never shows without its impact
        assert all(c["impact"] == [] for c in report["constants"] if not c["check"])
        assert bool(report["combined"]) == (report["status"] == "done")
    assert seen[-1]["progress"] == {"checked": 42, "total": 42}
