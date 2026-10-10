import datetime as dt
import re
import zipfile

import openpyxl
import pytest
from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.formula import ArrayFormula, DataTableFormula

from graph.parse import Graph, load


def _save(tmp_path, wb, name="model.xlsx"):
    path = tmp_path / name
    wb.save(path)
    return path


def _model(tmp_path):
    """Two data sheets (one with a space in its name) and an empty one."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Rev Build"
    ws["A1"] = "Growth"
    ws["B1"] = 0.08
    ws["B1"].number_format = "0.00%"
    ws["B1"].comment = Comment("from the 2023 board deck", "ethan")
    ws["B2"] = 1000
    ws["B3"] = "=B2*(1+B1)"
    ws["C5"] = "=SUM(B2:B3)"
    ws["D7"].comment = Comment("note on a blank cell", "ethan")

    out = wb.create_sheet("Output")
    out["C3"] = "='Rev Build'!B3*2"

    wb.create_sheet("Empty")
    return _save(tmp_path, wb)


def _set_cached(path, sheet_xml, cached):
    """Rewrite a saved xlsx so formula cells carry a cached <v>, as Excel leaves them. cached: {"B3": "1080"}."""
    with zipfile.ZipFile(path) as zf:
        parts = {info: zf.read(info.filename) for info in zf.infolist()}
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zf:
        for info, data in parts.items():
            if info.filename == sheet_xml:
                xml = data.decode()
                for coord, value in cached.items():
                    xml, count = re.subn(
                        rf'(<c r="{coord}"[^>]*>.*?</f>)<v ?/>', rf"\g<1><v>{value}</v>", xml
                    )
                    assert count == 1, coord
                data = xml.encode()
            zf.writestr(info.filename, data)


def test_sheets_and_keys(tmp_path):
    path = _model(tmp_path)
    wb = load(str(path))

    assert wb.path == path
    assert wb.sheets == ["Rev Build", "Output", "Empty"]
    assert set(wb.values) == {
        "Rev Build!A1",
        "Rev Build!B1",
        "Rev Build!B2",
        "Rev Build!B3",
        "Rev Build!C5",
        "Output!C3",
    }


def test_formulas_and_literals_are_separated(tmp_path):
    wb = load(_model(tmp_path))

    assert wb.formulas == {
        "Rev Build!B3": "=B2*(1+B1)",
        "Rev Build!C5": "=SUM(B2:B3)",
        "Output!C3": "='Rev Build'!B3*2",
    }
    assert wb.values["Rev Build!A1"] == "Growth"
    assert wb.values["Rev Build!B1"] == 0.08
    assert wb.values["Rev Build!B2"] == 1000
    assert type(wb.values["Rev Build!B2"]) is int


def test_cached_value_is_none_when_never_calculated(tmp_path):
    wb = load(_model(tmp_path))

    for key in wb.formulas:
        assert key in wb.values
        assert wb.values[key] is None


def test_cached_value_is_read_when_present(tmp_path):
    path = _model(tmp_path)
    _set_cached(path, "xl/worksheets/sheet1.xml", {"B3": "1080", "C5": "2080.5"})
    wb = load(path)

    assert wb.formulas["Rev Build!B3"] == "=B2*(1+B1)"
    assert wb.values["Rev Build!B3"] == 1080
    assert wb.values["Rev Build!C5"] == 2080.5
    assert wb.values["Output!C3"] is None


def test_string_that_looks_like_a_formula_is_a_literal(tmp_path):
    book = openpyxl.Workbook()
    ws = book.active
    ws["A1"] = "=not a formula"
    ws["A1"].data_type = "s"
    ws["A2"] = "=A3+1"
    wb = load(_save(tmp_path, book))

    assert wb.values["Sheet!A1"] == "=not a formula"
    assert wb.formulas == {"Sheet!A2": "=A3+1"}


def test_array_formula_text_and_data_table_skipped(tmp_path):
    book = openpyxl.Workbook()
    ws = book.active
    ws["A1"] = 1
    ws["A2"] = 2
    ws["B1"] = ArrayFormula("B1:B2", "=A1:A2*2")
    ws["C1"] = DataTableFormula(ref="C1:C2", r1="A1")
    path = _save(tmp_path, book)
    _set_cached(path, "xl/worksheets/sheet1.xml", {"B1": "2"})
    wb = load(path)

    assert wb.formulas == {"Sheet!B1": "=A1:A2*2"}
    assert wb.values["Sheet!B1"] == 2
    assert "Sheet!C1" not in wb.values
    assert "Sheet!C1" not in wb.formats


def test_number_formats(tmp_path):
    wb = load(_model(tmp_path))

    assert wb.formats["Rev Build!B1"] == "0.00%"
    assert wb.formats["Rev Build!B2"] == "General"
    assert wb.formats["Rev Build!B3"] == "General"
    assert set(wb.formats) == set(wb.values)


def test_comments_including_on_a_blank_cell(tmp_path):
    wb = load(_model(tmp_path))

    assert wb.comments == {
        "Rev Build!B1": "from the 2023 board deck",
        "Rev Build!D7": "note on a blank cell",
    }
    assert "Rev Build!D7" not in wb.values
    assert "Rev Build!D7" not in wb.formats


def test_bounds(tmp_path):
    wb = load(_model(tmp_path))

    assert wb.bounds == {
        "Rev Build": (1, 1, 7, 4),
        "Output": (3, 3, 3, 3),
        "Empty": (1, 1, 1, 1),
    }


def test_defined_names(tmp_path):
    book = openpyxl.Workbook()
    ws = book.active
    ws.title = "Assumptions"
    ws["B2"] = 0.08
    ws.print_area = "A1:B2"
    book.defined_names["Growth"] = DefinedName("Growth", attr_text="Assumptions!$B$2")
    book.defined_names["_xlnm.Criteria"] = DefinedName("_xlnm.Criteria", attr_text="Assumptions!$A$1")
    ws.defined_names["Local"] = DefinedName("Local", attr_text="Assumptions!$B$2")
    wb = load(_save(tmp_path, book))

    assert wb.names == {"GROWTH": "Assumptions!$B$2"}


def test_dates_and_booleans_kept_as_is(tmp_path):
    book = openpyxl.Workbook()
    ws = book.active
    ws["A1"] = dt.datetime(2024, 3, 31, 17, 30)
    ws["A2"] = dt.time(9, 15)
    ws["A3"] = True
    ws["A4"] = False
    ws["A5"] = dt.date(2024, 3, 31)
    wb = load(_save(tmp_path, book))

    assert wb.values["Sheet!A1"] == dt.datetime(2024, 3, 31, 17, 30)
    assert wb.values["Sheet!A2"] == dt.time(9, 15)
    assert wb.values["Sheet!A3"] is True
    assert wb.values["Sheet!A4"] is False
    # openpyxl reads every date-formatted number back as a datetime
    assert wb.values["Sheet!A5"] == dt.datetime(2024, 3, 31)


def test_chartsheets_are_skipped(tmp_path):
    book = openpyxl.Workbook()
    ws = book.active
    ws["A1"] = 1
    ws["A2"] = 2
    chart = BarChart()
    chart.add_data(Reference(ws, min_col=1, min_row=1, max_row=2))
    book.create_chartsheet("Chart").add_chart(chart)
    wb = load(_save(tmp_path, book))

    assert wb.sheets == ["Sheet"]
    assert list(wb.bounds) == ["Sheet"]


def test_invalid_file_raises(tmp_path):
    path = tmp_path / "broken.xlsx"
    path.write_bytes(b"not a workbook")

    with pytest.raises(zipfile.BadZipFile):
        load(path)


def test_context_on_a_loaded_workbook(tmp_path):
    wb = load(_model(tmp_path))
    assert wb.context("Rev Build!B1") == {
        "sheet": "Rev Build",
        "row_label": "Growth",
        "col_header": None,
        "number_format": "0.00%",
        "comment": "from the 2023 board deck",
    }


def _graph(edges):
    """edges: (precedent, dependent) pairs."""
    graph = Graph()
    for src, dst in edges:
        graph.dependents.setdefault(src, set()).add(dst)
        graph.precedents.setdefault(dst, set()).add(src)
    return graph


def test_downstream_chain():
    graph = _graph([("S!A1", "S!A2"), ("S!A2", "S!A3"), ("S!A3", "S!A4")])

    assert graph.downstream("S!A1") == {"S!A2", "S!A3", "S!A4"}
    assert graph.downstream("S!A3") == {"S!A4"}
    assert graph.downstream("S!A4") == set()


def test_downstream_diamond():
    graph = _graph([("S!A1", "S!B1"), ("S!A1", "S!C1"), ("S!B1", "S!D1"), ("S!C1", "S!D1")])

    assert graph.downstream("S!A1") == {"S!B1", "S!C1", "S!D1"}


def test_downstream_cycle_terminates_and_excludes_start():
    graph = _graph([("S!A1", "S!A2"), ("S!A2", "S!A3"), ("S!A3", "S!A1"), ("S!A3", "S!A4")])

    assert graph.downstream("S!A1") == {"S!A2", "S!A3", "S!A4"}
    assert graph.downstream("S!A2") == {"S!A1", "S!A3", "S!A4"}


def test_downstream_self_reference():
    graph = _graph([("S!A1", "S!A1")])

    assert graph.downstream("S!A1") == set()


def test_downstream_unknown_cell():
    graph = _graph([("S!A1", "S!A2")])

    assert graph.downstream("S!Z9") == set()
    assert graph.dependents == {"S!A1": {"S!A2"}}
