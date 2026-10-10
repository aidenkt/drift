import time
from pathlib import Path

import openpyxl
import pytest
from openpyxl.workbook.defined_name import DefinedName

from graph.parse import Literal, Workbook, build_graph, load
from graph.refs import split_key


def _wb(formulas, values=None, names=None, sheets=None):
    """Hand-built workbook. Formula cells are added to values; sheets default to those the keys mention."""
    values = {**{cell: None for cell in formulas}, **(values or {})}
    sheets = sheets or sorted({split_key(key)[0] for key in values})
    return Workbook(
        path=Path("x.xlsx"),
        sheets=sheets,
        formulas=dict(formulas),
        values=values,
        bounds={sheet: (1, 1, 1048576, 16384) for sheet in sheets},
        names=names or {},
    )


def _precedents(formula, values=None, names=None, sheets=None, cell="S!Z1"):
    return build_graph(_wb({cell: formula}, values, names, sheets)).precedents[cell]


def _literals(formula):
    return build_graph(_wb({"S!Z1": formula})).literals.get("S!Z1", [])


def _values(formula):
    return [literal.value for literal in _literals(formula)]


def _placed(formula):
    """(value, func, arg) per literal."""
    return [(literal.value, literal.func, literal.arg) for literal in _literals(formula)]


# edges


def test_same_sheet_edges():
    assert _precedents("=A1+B2", {"S!A1": 1, "S!B2": 2}) == {"S!A1", "S!B2"}


def test_cross_sheet_edge():
    assert _precedents("=Assumptions!B2*A1", {"Assumptions!B2": 0.08, "S!A1": 1}) == {
        "Assumptions!B2",
        "S!A1",
    }


def test_absolute_refs():
    values = {"S!A1": 1, "S!B2": 2, "S!C3": 3}
    assert _precedents("=$A$1+$B2+C$3", values) == {"S!A1", "S!B2", "S!C3"}


def test_lower_case_refs_and_sheet_name():
    assert _precedents("=assumptions!b2+a1", {"Assumptions!B2": 0.08, "S!A1": 1}) == {
        "Assumptions!B2",
        "S!A1",
    }


def test_quoted_sheet_name_with_space():
    assert _precedents("='Rev Build'!B3*2", {"Rev Build!B3": 5}) == {"Rev Build!B3"}


def test_sum_range_edges():
    values = {"S!A1": 1, "S!A2": 2, "S!A3": 3, "S!A4": 4}
    assert _precedents("=SUM(A1:A3)", values) == {"S!A1", "S!A2", "S!A3"}


def test_reversed_range():
    values = {"S!A1": 1, "S!B2": 2, "S!C3": 3}
    assert _precedents("=SUM(B2:A1)", values) == {"S!A1", "S!B2"}


def test_whole_column_and_whole_row_clipped_to_data():
    values = {"T!A1": 1, "T!A2": 2, "T!B2": 3, "T!C7": 4, "T!E1": 5}
    assert _precedents("=SUM(T!A:B)", values) == {"T!A1", "T!A2", "T!B2"}
    assert _precedents("=SUM(T!$1:$1)", values) == {"T!A1", "T!E1"}


def test_whole_column_ignores_inflated_bounds():
    """wb.bounds says the whole sheet is used; only the data is walked."""
    values = {f"T!A{row}": row for row in range(1, 4)}
    wb = _wb({"S!A1": "=SUM(T!A:XFD)"}, values)
    assert wb.bounds["T"] == (1, 1, 1048576, 16384)

    assert build_graph(wb).precedents["S!A1"] == {"T!A1", "T!A2", "T!A3"}


def test_stray_value_far_away_does_not_explode():
    values = {"T!A1": 1, "T!A2": 2, "T!B1": 3, "T!XFD1048576": "stray"}
    formulas = {
        "S!A1": "=SUM(T!A:A)",
        "S!A2": "=SUM(T!A:XFD)",
        "S!A3": "=SUM(T!1:1048576)",
        "S!A4": "=SUM(T!A1:XFD1048576)",
    }
    started = time.perf_counter()
    graph = build_graph(_wb(formulas, values))

    assert time.perf_counter() - started < 2
    assert graph.precedents["S!A1"] == {"T!A1", "T!A2"}
    for cell in ("S!A2", "S!A3", "S!A4"):
        assert graph.precedents[cell] == set(values)
    assert sum(len(cells) for cells in graph.precedents.values()) == 14


def test_blank_cell_references_dropped():
    values = {"S!A1": 1, "S!A3": 3}
    assert _precedents("=A1+B9+SUM(A2:A3)+Other!A1", values, sheets=["S", "Other"]) == {
        "S!A1",
        "S!A3",
    }


def test_sheet_with_no_values_is_still_a_known_sheet():
    assert _precedents("=SUM(Empty!A:A)+Empty!B2+SUM(Empty!A1:C3)", sheets=["S", "Empty"]) == set()


def test_unknown_sheet_gives_no_edge():
    assert _precedents("=Gone!A1+A1", {"S!A1": 1}) == {"S!A1"}


def test_defined_name_resolved():
    names = {"TAXRATE": "Assumptions!$B$15"}
    assert _precedents("=A1*TaxRate", {"S!A1": 1, "Assumptions!B15": 0.21}, names) == {
        "S!A1",
        "Assumptions!B15",
    }


def test_defined_name_with_leading_equals_and_range():
    values = {"Data!A1": 1, "Data!A2": 2, "Data!A9": 9}
    assert _precedents("=SUM(Sales)", values, {"SALES": "=Data!$A$1:$A$5"}) == {"Data!A1", "Data!A2"}


def test_defined_name_multi_area():
    values = {"Data!A1": 1, "Data!C3": 3, "A,B!D4": 4}
    names = {"PICKS": "Data!$A$1,Data!$C$3,'A,B'!$D$4"}
    assert _precedents("=SUM(Picks)", values, names) == set(values)


@pytest.mark.parametrize("value", ["0.08", "#REF!", '"text"', "#REF!$A$1", None])
def test_defined_name_without_cells_gives_no_edge(value):
    assert _precedents("=A1*Rate", {"S!A1": 1}, {"RATE": value}) == {"S!A1"}


def test_defined_name_the_tokenizer_reads_as_a_number():
    """float("Inf") succeeds, so the tokenizer calls a name spelled Inf a number."""
    graph = build_graph(_wb({"S!Z1": "=A1*(1+Inf)"}, {"S!A1": 1, "Assumptions!B4": 0.03}, {"INF": "Assumptions!$B$4"}))

    assert graph.precedents["S!Z1"] == {"S!A1", "Assumptions!B4"}
    assert [(lit.index, lit.value) for lit in graph.literals["S!Z1"]] == [(0, 1.0)]


def test_sheet_repeated_on_both_ends_of_a_range():
    values = {"Sheet1!A1": 1, "Sheet1!B2": 2, "Sheet1!C3": 3, "Rev Build!A1": 4, "Rev Build!A2": 5}
    assert _precedents("=SUM(Sheet1!A1:Sheet1!B2)", values) == {"Sheet1!A1", "Sheet1!B2"}
    assert _precedents("=SUM('Rev Build'!$A$1:'Rev Build'!$A$2)", values) == {
        "Rev Build!A1",
        "Rev Build!A2",
    }


def test_ref_error_tail_is_not_resolved_against_the_current_sheet():
    values = {"S!A1": 1, "S!B2": 2, "S!C3": 3}
    assert _precedents("=#REF!A1+B2", values) == {"S!B2"}
    assert _precedents("=SUM(#REF!A1:C3,B2)", values) == {"S!B2"}
    assert _precedents("=#REF!+A1", values) == {"S!A1"}
    assert _precedents("=#N/A+A1", values) == {"S!A1"}


def test_external_workbook_refs_ignored():
    values = {"S!A1": 1, "Sheet1!A1": 2}
    formula = "=[Budget.xlsx]Sheet1!A1+'[Budget.xlsx]Sheet1'!A1+[1]Sheet1!A1+A1"
    assert _precedents(formula, values) == {"S!A1"}


def test_indirect_and_offset_take_their_literal_references():
    values = {"S!A1": 1, "S!B5": 2}
    assert _precedents('=INDIRECT("B5")+OFFSET(A1,1,2)', values) == {"S!A1"}


def test_range_ending_in_a_function():
    """The tokenizer glues "A1:INDEX(" into one function token."""
    graph = build_graph(_wb({"S!Z1": "=SUM(A1:INDEX(B:B,5))"}, {"S!A1": 1, "S!B3": 2}))

    assert graph.precedents["S!Z1"] == {"S!A1", "S!B3"}
    assert graph.literals["S!Z1"] == [Literal(cell="S!Z1", index=0, value=5.0, func="INDEX", arg=1)]


def test_formula_without_references_has_empty_precedents():
    graph = build_graph(_wb({"S!A1": "=1+2", "S!A2": "=B9"}))

    assert graph.precedents == {"S!A1": set(), "S!A2": set()}
    assert graph.dependents == {}


def test_dependents_is_the_exact_reverse_of_precedents():
    formulas = {
        "S!A3": "=A1+A2",
        "S!A4": "=SUM(A1:A3)*Rate",
        "S!A5": "=A4+Other!B1+A9",
        "Other!B1": "=S!A1",
    }
    graph = build_graph(_wb(formulas, {"S!A1": 1, "S!A2": 2, "Other!C1": 0.5}, {"RATE": "Other!$C$1"}))

    assert graph.precedents == {
        "S!A3": {"S!A1", "S!A2"},
        "S!A4": {"S!A1", "S!A2", "S!A3", "Other!C1"},
        "S!A5": {"S!A4", "Other!B1"},
        "Other!B1": {"S!A1"},
    }
    reverse = {}
    for cell, precedents in graph.precedents.items():
        for precedent in precedents:
            reverse.setdefault(precedent, set()).add(cell)
    assert graph.dependents == reverse
    assert graph.downstream("S!A1") == {"S!A3", "S!A4", "S!A5", "Other!B1"}


# literals


def test_literal_values():
    assert _values("=B5*1.08") == [1.08]
    assert _values("=-5") == [-5.0]
    assert _values("=B5*-2") == [-2.0]
    assert _values("=A1-5") == [5.0]
    assert _values("=+5") == [5.0]
    assert _values("=5%") == [0.05]
    assert _values("=-8%") == [-0.08]
    assert _values("=1E+3+1.5e-3+.5") == [1000.0, 0.0015, 0.5]


def test_whitespace_between_sign_number_and_percent():
    assert _values("= A1 - 5 + - 3 %") == [5.0, -0.03]


def test_prefix_minus_on_a_parenthesis_does_not_negate_the_number_inside():
    assert _values("=-(5)") == [5.0]


def test_literal_fields_and_index_numbering():
    assert _literals("=B5*1.08+A1*3") == [
        Literal(cell="S!Z1", index=0, value=1.08, func=None, arg=None),
        Literal(cell="S!Z1", index=1, value=3.0, func=None, arg=None),
    ]


def test_cells_without_number_literals_have_no_entry():
    graph = build_graph(_wb({"S!A1": "=B1+C1", "S!A2": '="x 7"&B1', "S!A3": "=TRUE", "S!A4": "=B1*2"}))

    assert list(graph.literals) == ["S!A4"]


def test_func_and_arg():
    assert _placed("=ROUND(B2*1.08,2)") == [(1.08, "ROUND", 0), (2.0, "ROUND", 1)]


def test_func_and_arg_nested():
    assert _placed("=IF(A1>0,ROUND(B1*1.5,2),3)") == [
        (0.0, "IF", 0),
        (1.5, "ROUND", 0),
        (2.0, "ROUND", 1),
        (3.0, "IF", 2),
    ]


def test_top_level_literals_have_no_func():
    assert _placed("=(B5+2)*1.08+SUM(A1:A3)-4") == [
        (2.0, None, None),
        (1.08, None, None),
        (4.0, None, None),
    ]


def test_array_constant_does_not_advance_the_outer_arg():
    assert _placed("=SUM({1,2;3,4},7)") == [
        (1.0, "SUM", 0),
        (2.0, "SUM", 0),
        (3.0, "SUM", 0),
        (4.0, "SUM", 0),
        (7.0, "SUM", 1),
    ]
    assert _placed("=A1*{1,2}") == [(1.0, None, None), (2.0, None, None)]


def test_union_parentheses_do_not_advance_the_outer_arg():
    assert _placed("=SUM((A1,B2),3)") == [(3.0, "SUM", 1)]
    assert _placed("=ROUND((B1+1)*(2+C1),4)") == [(1.0, "ROUND", 0), (2.0, "ROUND", 0), (4.0, "ROUND", 1)]


def test_empty_arguments_are_counted():
    assert _placed("=IF(A1,,5)") == [(5.0, "IF", 2)]


def test_func_name_is_upper_case_without_storage_prefix():
    assert _placed("=round(A1,2)") == [(2.0, "ROUND", 1)]
    assert _placed("=_xlfn.XLOOKUP(A1,B:B,C:C,0)") == [(0.0, "XLOOKUP", 3)]
    assert _placed("=_xlfn._xlws.FILTER(A1:A3,B1:B3>2)") == [(2.0, "FILTER", 1)]


def test_prefix_and_percent_inside_a_function():
    assert _placed("=ROUND(-B2*1.08,2)+5%-VLOOKUP(A1,T!A:C,3,0)*1E+3") == [
        (1.08, "ROUND", 0),
        (2.0, "ROUND", 1),
        (0.05, None, None),
        (3.0, "VLOOKUP", 2),
        (0.0, "VLOOKUP", 3),
        (1000.0, None, None),
    ]
    assert _placed("=MAX(-1,A1*-50%)") == [(-1.0, "MAX", 0), (-0.5, "MAX", 1)]


def test_numbers_inside_text_are_not_literals():
    assert _values('="x 7"&A1&TEXT(B1,"0.00")') == []


def test_non_finite_numbers_are_skipped_but_still_numbered():
    assert [(lit.index, lit.value) for lit in _literals("=1E999+2")] == [(1, 2.0)]


# cycles


def test_two_cell_cycle():
    graph = build_graph(_wb({"S!A1": "=B1+1", "S!B1": "=A1*2", "S!C1": "=A1"}))

    assert graph.cycles == [["S!A1", "S!B1"]]


def test_self_reference_is_a_cycle():
    graph = build_graph(_wb({"S!A1": "=A1+1", "S!B1": "=SUM(B1:B3)", "S!C1": "=A1"}, {"S!B2": 1}))

    assert graph.precedents["S!A1"] == {"S!A1"}
    assert graph.dependents["S!A1"] == {"S!A1", "S!C1"}
    assert graph.cycles == [["S!A1"], ["S!B1"]]


def test_no_cycles_in_a_dag():
    formulas = {"S!B1": "=A1", "S!C1": "=A1+B1", "S!D1": "=B1+C1", "S!E1": "=SUM(A1:D1)"}
    graph = build_graph(_wb(formulas, {"S!A1": 1}))

    assert graph.cycles == []


def test_cycles_are_sorted_and_separate():
    formulas = {
        "T!C1": "=T!A1",
        "T!A1": "=T!B1",
        "T!B1": "=T!C1+S!A1",
        "S!B1": "=S!A1",
        "S!A1": "=S!B1",
        "S!D1": "=S!A1+T!A1",
    }
    graph = build_graph(_wb(formulas))

    assert graph.cycles == [["S!A1", "S!B1"], ["T!A1", "T!B1", "T!C1"]]


def test_cycle_through_a_range_and_a_defined_name():
    formulas = {"S!A5": "=SUM(A1:A4)", "S!A2": "=Total*0.5"}
    graph = build_graph(_wb(formulas, {"S!A1": 1}, {"TOTAL": "S!$A$5"}))

    assert graph.cycles == [["S!A2", "S!A5"]]


def test_long_chain_does_not_hit_the_recursion_limit():
    formulas = {f"S!A{row}": f"=A{row - 1}+1" for row in range(2, 5001)}
    graph = build_graph(_wb(formulas, {"S!A1": 1}))

    assert graph.cycles == []
    assert graph.precedents["S!A5000"] == {"S!A4999"}
    assert len(graph.downstream("S!A1")) == 4999


def test_long_cycle_does_not_hit_the_recursion_limit():
    formulas = {f"S!A{row}": f"=A{row - 1}+1" for row in range(2, 5001)}
    formulas["S!A1"] = "=A5000"
    graph = build_graph(_wb(formulas))

    assert graph.cycles == [sorted(formulas)]


# robustness


@pytest.mark.parametrize(
    "formula",
    [
        "=A1)+(",
        '="abc',
        "=SUM(1}",
        "=)",
        "=SUM(A1:A3)#",
        "=SUM(A1",
        "={1,2",
        "=",
        "",
        "A1+1",
        "=1E+",
        "=1..2",
        "=A1:A" + "9" * 5000,
        "=Sheet1!A1:Sheet2!B2",
        "=Sheet1:Sheet3!A1",
        "=Table1[Col]+@A1",
        "='" + "x" * 3000,
        None,
    ],
)
def test_malformed_formula_does_not_raise(formula):
    graph = build_graph(_wb({"S!Z1": formula, "S!Z2": "=A1*2"}, {"S!A1": 1}))

    assert "S!Z1" in graph.precedents
    assert graph.precedents["S!Z2"] == {"S!A1"}
    assert graph.literals["S!Z2"] == [Literal(cell="S!Z2", index=0, value=2.0)]


def test_formula_the_tokenizer_rejects_has_no_precedents_or_literals():
    graph = build_graph(_wb({"S!Z1": "=A1*2)+("}, {"S!A1": 1}))

    assert graph.precedents == {"S!Z1": set()}
    assert graph.literals == {}
    assert graph.cycles == []


def test_empty_workbook():
    graph = build_graph(Workbook(path=Path("x.xlsx"), sheets=[]))

    assert (graph.precedents, graph.dependents, graph.literals, graph.cycles) == ({}, {}, {}, [])


def test_shared_whole_column_lookups_build_quickly():
    """1,000 formulas reading the same two 1,000-row columns: each range is expanded once."""
    values = {f"Data!{col}{row}": row for row in range(1, 1001) for col in "AB"}
    formulas = {f"S!B{row}": f"=SUMIF(Data!$A:$A,A{row},Data!$B:$B)*1.1" for row in range(1, 1001)}
    started = time.perf_counter()
    graph = build_graph(_wb(formulas, values))

    assert time.perf_counter() - started < 5
    assert len(graph.precedents["S!B7"]) == 2000
    assert len(graph.dependents["Data!A1"]) == 1000


# through a real file


def _save(tmp_path, book):
    path = tmp_path / "model.xlsx"
    book.save(path)
    return path


def test_real_file_edges_literals_and_names(tmp_path):
    book = openpyxl.Workbook()
    assumptions = book.active
    assumptions.title = "Assumptions"
    assumptions["A1"] = "Growth"
    assumptions["B1"] = 0.08
    assumptions["B2"] = 0.21
    book.defined_names["TaxRate"] = DefinedName("TaxRate", attr_text="Assumptions!$B$2")
    book.defined_names["Floor"] = DefinedName("Floor", attr_text="100")

    rev = book.create_sheet("Rev Build")
    rev["A1"] = "Revenue"
    rev["B1"] = 1000
    rev["C1"] = "=ROUND(B1*(1+assumptions!$B$1),2)"
    rev["D1"] = "=C1*1.08-5%"
    rev["E1"] = "=SUM(B1:D1)"
    rev["F1"] = "=E1*(1-TaxRate)+Floor"

    out = book.create_sheet("Output")
    out["A1"] = "='Rev Build'!F1/1E+3"
    out["A2"] = "=SUM('Rev Build'!B:F)+Z99"

    graph = build_graph(load(_save(tmp_path, book)))

    assert graph.precedents == {
        "Rev Build!C1": {"Rev Build!B1", "Assumptions!B1"},
        "Rev Build!D1": {"Rev Build!C1"},
        "Rev Build!E1": {"Rev Build!B1", "Rev Build!C1", "Rev Build!D1"},
        "Rev Build!F1": {"Rev Build!E1", "Assumptions!B2"},
        "Output!A1": {"Rev Build!F1"},
        "Output!A2": {f"Rev Build!{col}1" for col in "BCDEF"},
    }
    assert graph.dependents["Assumptions!B1"] == {"Rev Build!C1"}
    assert graph.dependents["Rev Build!F1"] == {"Output!A1", "Output!A2"}
    assert graph.downstream("Assumptions!B1") == {
        "Rev Build!C1",
        "Rev Build!D1",
        "Rev Build!E1",
        "Rev Build!F1",
        "Output!A1",
        "Output!A2",
    }
    assert graph.literals == {
        "Rev Build!C1": [
            Literal(cell="Rev Build!C1", index=0, value=1.0, func="ROUND", arg=0),
            Literal(cell="Rev Build!C1", index=1, value=2.0, func="ROUND", arg=1),
        ],
        "Rev Build!D1": [
            Literal(cell="Rev Build!D1", index=0, value=1.08),
            Literal(cell="Rev Build!D1", index=1, value=0.05),
        ],
        "Rev Build!F1": [Literal(cell="Rev Build!F1", index=0, value=1.0)],
        "Output!A1": [Literal(cell="Output!A1", index=0, value=1000.0)],
    }
    assert graph.cycles == []


def test_real_file_with_style_only_cell_far_away_and_a_cycle(tmp_path):
    book = openpyxl.Workbook()
    data = book.active
    data.title = "Data"
    for row in range(1, 6):
        data[f"A{row}"] = row
    data["XFD1048576"].number_format = "0.00"
    data["C1"] = "=C2+1"
    data["C2"] = "=C1*2"

    calc = book.create_sheet("Calc")
    calc["A1"] = "=SUM(Data!A:A)"
    calc["A2"] = "=SUM(Data!1:1048576)"
    calc["A3"] = "=A3+Data!A1"
    book.create_sheet("Empty")

    wb = load(_save(tmp_path, book))
    assert wb.bounds["Data"] == (1, 1, 1048576, 16384)

    started = time.perf_counter()
    graph = build_graph(wb)

    assert time.perf_counter() - started < 2
    assert graph.precedents["Calc!A1"] == {f"Data!A{row}" for row in range(1, 6)}
    assert graph.precedents["Calc!A2"] == {f"Data!A{row}" for row in range(1, 6)} | {"Data!C1", "Data!C2"}
    assert graph.precedents["Calc!A3"] == {"Calc!A3", "Data!A1"}
    assert sum(len(cells) for cells in graph.precedents.values()) == 16
    assert graph.cycles == [["Calc!A3"], ["Data!C1", "Data!C2"]]
