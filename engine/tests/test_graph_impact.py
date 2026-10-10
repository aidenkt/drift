import math

import openpyxl
import pytest
from openpyxl.worksheet.formula import ArrayFormula, DataTableFormula

from graph import impact, recalc
from graph.parse import Workbook, build_graph, load

EBITDA, NET, NPV, ADJUSTED, LOOKUP, RATIO, COVERAGE = (f"P&L!B{row}" for row in range(2, 9))
MAIN = [EBITDA, NET, NPV]
COST = ["Rev Build!B2", "Rev Build!C2", "Rev Build!D2"]

# the model below, worked by hand
BASELINE = {EBITDA: 2700.0, NET: 1975.0, NPV: 1580.0, ADJUSTED: 1895.0, LOOKUP: 13500.0, COVERAGE: 13.5}


def _write_model(path):
    """Revenue 1000 / 1500 / 2000, cost 40% of each, EBITDA 2700, net income 1975, NPV 1580."""
    book = openpyxl.Workbook()
    inputs = book.active
    inputs.title = "Inputs"
    rows = [("Price", 10), ("Volume", 100), ("Tariff", 0), ("Step", 500), ("Discount rate", 0.25), ("One-off", 200)]
    for row, (label, value) in enumerate(rows, start=1):
        inputs.cell(row, 1, label)
        inputs.cell(row, 2, value)

    rev = book.create_sheet("Rev Build")
    rev["A1"] = "Revenue"
    rev["B1"] = "=Inputs!B1*Inputs!B2"
    rev["C1"] = "=B1+Inputs!$B$4"
    rev["D1"] = "=C1+Inputs!$B$4"
    rev["A2"] = "Cost"
    for col in "BCD":
        rev[f"{col}2"] = f"={col}1*0.4"

    pnl = book.create_sheet("P&L")
    formulas = [
        ("Tariff cost", "='Rev Build'!B1*Inputs!B3"),
        ("EBITDA", "=SUM('Rev Build'!B1:D1)-SUM('Rev Build'!B2:D2)-B1"),
        ("Net income", "=B2*(1-25%)-50"),
        ("NPV", "=B3/(1+Inputs!B5)"),
        ("Adjusted", "=B3+ - 40 %*Inputs!B6"),
        ("Lookup", "=SUM({2,3})*B2"),
        ("Ratio", "=B2/Inputs!B3"),
        ("Coverage", "=B2/Inputs!B6"),
    ]
    for row, (label, formula) in enumerate(formulas, start=1):
        pnl.cell(row, 1, label)
        pnl.cell(row, 2, formula)
    book.save(path)


def _load(path, cached=None):
    """(wb, graph); cached stands in for the results Excel would have stored in the file."""
    wb = load(path)
    wb.values.update(cached or {})
    return wb, build_graph(wb)


def _outputs(*cells, value=0.0):
    return [{"cell": cell, "label": cell, "value": value} for cell in cells or BASELINE]


def _input(cell, value, live, outputs=MAIN):
    return {"cell": cell, "value": value, "embedded": False, "outputs": outputs, "check": {"live": live}}


def _literal(cells, index, value, live, outputs=MAIN):
    constant = _input(cells[0], value, live, outputs) | {"embedded": True, "literal_index": index}
    if len(cells) > 1:
        constant["cell"] = f"{cells[0]}:{cells[-1].rpartition('!')[2]}"
        constant["cells"] = cells
    return constant


def _by_output(result):
    assert all(type(item["delta"]) is float for item in result)
    assert len({item["output"] for item in result}) == len(result)
    return {item["output"]: item["delta"] for item in result}


def _deltas(model, constant, outputs=None):
    return _by_output(impact.deltas(*model, constant, outputs or _outputs()))


@pytest.fixture(scope="module")
def model(tmp_path_factory):
    path = tmp_path_factory.mktemp("model") / "model.xlsx"
    _write_model(path)
    return _load(path, BASELINE)


@pytest.fixture(scope="module")
def uncalculated(tmp_path_factory):
    """The same model as openpyxl saves it: no cached results at all."""
    path = tmp_path_factory.mktemp("uncalculated") / "model.xlsx"
    _write_model(path)
    return _load(path)


@pytest.fixture(scope="module")
def named(tmp_path_factory):
    """Sheet names with a space, an apostrophe and a percent sign, next to sheets that already have
    the names the recalculation would pick for its scratch sheet and for the renamed apostrophe sheet."""
    book = openpyxl.Workbook()
    rev = book.active
    rev.title = "Rev Build"
    rev["B1"] = 4
    rev["B2"] = "=_drift!A1+B1"
    rev["B3"] = "='Bob_s Summary'!A1+'Bob''s Summary'!B2"
    rev["B4"] = "=SUM('Margin %'!B1:B2)"
    summary = book.create_sheet("Bob's Summary")
    summary["B1"] = "='Rev Build'!B1*2.5"
    summary["B2"] = "=B1*3"
    margin = book.create_sheet("Margin %")
    margin["B1"] = "='Bob''s Summary'!B2/100"
    margin["B2"] = "=SUM('Bob''s Summary'!B1:B2)*B1"
    book.create_sheet("_drift")["A1"] = 7
    book.create_sheet("Bob_s Summary")["A1"] = 1
    path = tmp_path_factory.mktemp("named") / "named.xlsx"
    book.save(path)
    cached = {"Rev Build!B2": 11, "Rev Build!B3": 31, "Rev Build!B4": 12.3, "Bob's Summary!B2": 30, "Margin %!B2": 12}
    return _load(path, cached)


def _tiny(tmp_path, cached=None):
    """B1 = 2, B2 = B1 * 2.5."""
    book = openpyxl.Workbook()
    ws = book.active
    ws.title = "S"
    ws["B1"] = 2
    ws["B2"] = "=B1*2.5"
    path = tmp_path / "tiny.xlsx"
    book.save(path)
    return _load(path, cached)


# recalc


def test_hoist_moves_literals_to_scratch_cells(tmp_path):
    book = openpyxl.Workbook()
    ws = book.active
    ws.title = "S"
    ws["B2"] = 10
    ws["B3"] = "=B2*(1+5%)"
    ws["B4"] = "=-3*B2-2"
    ws["B5"] = "=SUM({1,2;3,4})*2.5"
    ws["B6"] = ArrayFormula("B6", "=B2*7")
    ws["B7"] = 5
    ws["B8"] = "= - 40 % * B2"
    ws["B9"] = "=B2*1.5"
    ws["B10"] = DataTableFormula("B10:B11", r1="B2")
    path = tmp_path / "source.xlsx"
    book.save(path)
    dest = tmp_path / "copy.xlsx"

    literals = [("S!B3", 1), ("S!B4", 0), ("S!B4", 1), ("S!B5", 0), ("S!B5", 4), ("S!B6", 0), ("S!B7", 0)]
    literals += [("S!B8", 0), ("Nope!A1", 0), ("S!B3", 9)]
    hoisted = recalc.hoist(load(path), literals, dest)

    assert hoisted == {
        ("S!B3", 1): "_drift!A1",
        ("S!B4", 0): "_drift!A2",
        ("S!B4", 1): "_drift!A3",
        ("S!B5", 4): "_drift!A4",
        ("S!B8", 0): "_drift!A5",
    }
    copy = openpyxl.load_workbook(dest)
    assert [cell.value for cell in copy["_drift"]["A"]] == [0.05, -3, 2, 2.5, -0.4]
    assert copy["S"]["B3"].value == "=B2*(1+_drift!$A$1)"
    assert copy["S"]["B4"].value == "=_drift!$A$2*B2-_drift!$A$3"
    assert copy["S"]["B5"].value == "=SUM({1,2;3,4})*_drift!$A$4"
    assert copy["S"]["B8"].value == "= _drift!$A$5 * B2"
    assert copy["S"]["B9"].value == "=B2*1.5"
    assert copy["S"]["B6"].value.text == "=B2*7"
    assert copy["S"]["B10"].value is None
    source = openpyxl.load_workbook(path)["S"]
    assert source["B3"].value == "=B2*(1+5%)"
    assert isinstance(source["B10"].value, DataTableFormula)


@pytest.fixture(scope="module")
def engine(tmp_path_factory):
    book = openpyxl.Workbook()
    ws = book.active
    ws.title = "S"
    ws["B1"] = 8
    ws["B2"] = "=B1*-25%"
    ws["B3"] = "=1/B1"
    ws["B4"] = '="abc"'
    ws["B5"] = "=B1>0"
    ws["B6"] = "=B1/0"
    ws["B7"] = "label"
    path = tmp_path_factory.mktemp("engine") / "engine.xlsx"
    book.save(path)
    return recalc.compile(load(path), [("S!B2", 0)])


def test_calculate(engine):
    cells = ["S!B1", "S!B2", "S!B3", "S!B4", "S!B5", "S!B6", "S!B7", "S!Z9", "Nope!A1"]
    scratch = engine.scratch["S!B2", 0]

    assert recalc.calculate(engine, {}, cells) == {
        "S!B1": 8.0,
        "S!B2": -2.0,
        "S!B3": 0.125,
        "S!B4": None,
        "S!B5": None,
        "S!B6": None,
        "S!B7": None,
        "S!Z9": None,
        "Nope!A1": None,
    }
    assert recalc.calculate(engine, {"S!B1": 4, scratch: 0.5}, ["S!B2", "S!B3"]) == {"S!B2": 2.0, "S!B3": 0.25}
    assert recalc.calculate(engine, {"S!B1": 0}, ["S!B2", "S!B3"]) == {"S!B2": 0.0, "S!B3": None}
    assert recalc.calculate(engine, {}, ["S!B2", "S!B3"]) == {"S!B2": -2.0, "S!B3": 0.125}
    assert all(type(value) is float for value in recalc.calculate(engine, {}, ["S!B1", "S!B2"]).values())


def test_calculate_refuses_an_override_of_a_cell_it_does_not_have(engine):
    assert recalc.calculate(engine, {"S!Z9": 1.0}, ["S!B2"]) == {"S!B2": None}
    assert recalc.calculate(engine, {"Nope!A1": 1.0}, ["S!B2"]) == {"S!B2": None}
    assert recalc.calculate(engine, {}, []) == {}


def test_compile_refuses_big_workbooks(tmp_path, monkeypatch):
    wb, _ = _tiny(tmp_path)
    monkeypatch.setattr(recalc, "MAX_FORMULAS", 0)
    assert recalc.compile(wb, []) is None

    monkeypatch.undo()
    book = openpyxl.Workbook()
    book.active["A1"] = 1
    book.active["B1"] = "=A1*2"
    book.active["Z40000"] = 1
    path = tmp_path / "sparse.xlsx"
    book.save(path)
    assert recalc.compile(load(path), []) is None


def test_compile_returns_none_when_it_fails(tmp_path):
    assert recalc.compile(Workbook(path=tmp_path / "missing.xlsx", sheets=[]), []) is None


# impact


def test_cell_input(model):
    # price 10 -> 12: revenue 1200 + 1700 + 2200 = 5100, cost 2040, EBITDA 3060, net 2245, NPV 1796
    assert _deltas(model, _input("Inputs!B1", 10, 12)) == pytest.approx({EBITDA: 360, NET: 270, NPV: 216})


def test_embedded_literal_in_one_cell(model):
    # the 50 in =B2*(1-25%)-50 -> 80
    assert _deltas(model, _literal([NET], 2, 50, 80)) == pytest.approx({EBITDA: 0, NET: -30, NPV: -24})


def test_two_literals_in_one_formula_move_separately(model):
    # the 25% in =B2*(1-25%)-50 -> 30%: net 2700 * 0.7 - 50 = 1840, NPV 1472
    assert _deltas(model, _literal([NET], 1, 0.25, 0.30)) == pytest.approx({EBITDA: 0, NET: -135, NPV: -108})


def test_embedded_group_moves_every_cell(model):
    # cost 40% -> 50% of 4500 revenue, against 40% -> 50% of the first year's 1000 only
    group = _deltas(model, _literal(COST, 0, 0.4, 0.5))
    first = _deltas(model, _literal(COST[:1], 0, 0.4, 0.5))

    assert group == pytest.approx({EBITDA: -450, NET: -337.5, NPV: -270})
    assert first == pytest.approx({EBITDA: -100, NET: -75, NPV: -60})


def test_prefix_minus_and_percent_are_part_of_the_literal(model):
    # =B3+ - 40 %*Inputs!B6 holds -0.4; at -0.25 it is 1975 - 50 = 1925
    constant = _literal([ADJUSTED], 0, -0.4, -0.25, outputs=[ADJUSTED])
    assert _deltas(model, constant) == pytest.approx({ADJUSTED: 30})


def test_hoisting_leaves_the_baseline_alone(uncalculated):
    outputs = _outputs()
    impact.fill_outputs(*uncalculated, outputs)
    assert {output["cell"]: output["value"] for output in outputs} == pytest.approx(BASELINE)


def test_literal_in_an_array_constant_has_no_impact(model):
    constant = _literal([LOOKUP], 0, 2, 4, outputs=[LOOKUP])
    assert impact.deltas(*model, constant, _outputs()) == []


@pytest.mark.parametrize(
    "check", [None, {}, {"live": None}, {"live": True}, {"live": math.nan}, {"live": math.inf}, {"live": "12"}]
)
def test_no_live_value_no_impact(model, check):
    constant = _input("Inputs!B1", 10, 12) | {"check": check}
    assert impact.deltas(*model, constant, _outputs()) == []
    assert impact.combined(*model, [constant], _outputs()) == []


@pytest.mark.parametrize("live, plausible", [(1001, False), (-1001, False), (0.099, False), (1000, True), (0.1, True)])
def test_live_value_in_other_units_is_skipped(model, live, plausible):
    assert bool(impact.deltas(*model, _input("Inputs!B1", 10, live), _outputs())) is plausible


def test_zero_constant_can_take_a_live_value(model):
    # tariff 0 -> 50% of the first year's 1000 revenue
    assert _deltas(model, _input("Inputs!B3", 0, 0.5)) == pytest.approx({EBITDA: -500, NET: -375, NPV: -300})


def test_only_outputs_the_constant_reaches_in_output_order(model):
    constant = _input("Inputs!B1", 10, 12, outputs=[NPV, "P&L!B99", EBITDA])
    result = impact.deltas(*model, constant, _outputs())

    assert [item["output"] for item in result] == [EBITDA, NPV]
    assert impact.deltas(*model, constant, _outputs(NET)) == []
    assert impact.deltas(*model, _input("Inputs!B1", 10, 12, outputs=[]), _outputs()) == []


def test_output_that_disagrees_with_excel_gets_no_impact(model, monkeypatch):
    monkeypatch.setitem(model[0].values, EBITDA, 2500.0)
    monkeypatch.setitem(model[0].values, NET, "n/a")

    assert _deltas(model, _input("Inputs!B1", 10, 12)) == pytest.approx({NPV: 216})
    assert EBITDA not in _by_output(impact.combined(*model, [_input("Inputs!B1", 10, 12)], _outputs()))


def test_cached_value_may_be_rounded(model, monkeypatch):
    monkeypatch.setitem(model[0].values, EBITDA, 2700.001)
    monkeypatch.setitem(model[0].values, NET, 1975.01)

    assert _deltas(model, _input("Inputs!B1", 10, 12)) == pytest.approx({EBITDA: 360, NPV: 216})


def test_output_that_is_an_error_gets_no_impact(model):
    # ratio divides by the tariff: #DIV/0! in the workbook, a number once the tariff is 50%
    tariff = _input("Inputs!B3", 0, 0.5, outputs=[EBITDA, RATIO])
    assert _deltas(model, tariff, _outputs(EBITDA, RATIO)) == pytest.approx({EBITDA: -500})

    # coverage divides by the one-off: a number in the workbook, #DIV/0! once the one-off is 0
    one_off = _input("Inputs!B6", 200, 0, outputs=[ADJUSTED, COVERAGE])
    assert _deltas(model, one_off, _outputs(ADJUSTED, COVERAGE)) == pytest.approx({ADJUSTED: 80})


def test_constant_the_model_cannot_override_has_no_impact(model):
    assert impact.deltas(*model, _input("Inputs!Z99", 10, 12), _outputs()) == []
    assert impact.deltas(*model, _literal([NET], 7, 50, 80), _outputs()) == []
    assert impact.deltas(*model, _literal([COST[0], "Rev Build!E2"], 0, 0.4, 0.5), _outputs()) == []


def test_calls_do_not_affect_each_other(model):
    price = _deltas(model, _input("Inputs!B1", 10, 12))
    # price 10 -> 20: revenue 7500, EBITDA 4500
    assert _deltas(model, _input("Inputs!B1", 10, 20)) == pytest.approx({EBITDA: 1800, NET: 1350, NPV: 1080})
    assert _deltas(model, _literal(COST, 0, 0.4, 0.5)) == pytest.approx({EBITDA: -450, NET: -337.5, NPV: -270})
    assert _deltas(model, _input("Inputs!B1", 10, 12)) == price == pytest.approx({EBITDA: 360, NET: 270, NPV: 216})


def test_combined_is_the_joint_effect(model):
    # price 12 and cost 50% together: revenue 5100, cost 2550, EBITDA 2550, net 1862.5, NPV 1490
    price = _input("Inputs!B1", 10, 12)
    cost = _literal(COST, 0, 0.4, 0.5, outputs=[EBITDA])
    ignored = [_input("Inputs!B4", 500, None), _input("Inputs!B5", 0.25, 80), _literal([LOOKUP], 0, 2, 4, [LOOKUP])]

    result = _by_output(impact.combined(*model, [price, cost, *ignored], _outputs()))

    assert result == pytest.approx(
        {EBITDA: -150, NET: -112.5, NPV: -90, ADJUSTED: -112.5, LOOKUP: -750, COVERAGE: -0.75}
    )
    assert list(result) == list(BASELINE)
    assert result[EBITDA] != pytest.approx(_deltas(model, price)[EBITDA] + _deltas(model, cost)[EBITDA])


def test_combined_without_a_qualifying_constant(model):
    assert impact.combined(*model, [], _outputs()) == []
    assert impact.combined(*model, [_input("Inputs!B1", 10, None)], _outputs()) == []
    # a live value, but no output to move
    assert impact.combined(*model, [_input("Inputs!B1", 10, 12, outputs=[])], _outputs()) == []


def test_fill_outputs_only_where_excel_left_no_value(model, uncalculated):
    cached = _outputs(value=-1.0)
    impact.fill_outputs(*model, cached)
    assert [output["value"] for output in cached] == [-1.0] * len(cached)

    filled = _outputs(EBITDA, RATIO, NPV, value=-1.0)
    impact.fill_outputs(*uncalculated, filled)
    assert [output["value"] for output in filled] == pytest.approx([2700, -1.0, 1580])
    assert type(filled[0]["value"]) is float


def test_uncalculated_workbook_still_gets_impacts(uncalculated):
    assert _deltas(uncalculated, _input("Inputs!B1", 10, 12)) == pytest.approx({EBITDA: 360, NET: 270, NPV: 216})


def test_sheet_names_with_space_apostrophe_and_percent(named):
    summary, npv, total, both, own = "Bob's Summary!B2", "Margin %!B2", "Rev Build!B3", "Rev Build!B4", "Rev Build!B2"
    outputs = _outputs(summary, npv, total, both, own)
    every = [summary, npv, total, both, own]

    # units 4 -> 6: revenue 15, EBITDA 45, margin 0.45, NPV 60 * 0.45 = 27
    units = _deltas(named, _input("Rev Build!B1", 4, 6, every), outputs)
    assert units == pytest.approx({summary: 15, npv: 15, total: 15, both: 15.15, own: 2})
    # the 2.5 in ='Rev Build'!B1*2.5 -> 5: revenue 20, EBITDA 60, margin 0.6, NPV 80 * 0.6 = 48
    price = _deltas(named, _literal(["Bob's Summary!B1"], 0, 2.5, 5, every), outputs)
    assert price == pytest.approx({summary: 30, npv: 36, total: 30, both: 36.3, own: 0})
    # the 100 in ='Bob''s Summary'!B2/100 -> 50: margin 0.6, NPV 40 * 0.6 = 24
    scale = _deltas(named, _literal(["Margin %!B1"], 0, 100, 50, [npv, both]), outputs)
    assert scale == pytest.approx({npv: 12, both: 12.3})
    # a cell on the workbook's own _drift sheet is still that cell: 7 -> 9
    assert _deltas(named, _input("_drift!A1", 7, 9, [own]), outputs) == pytest.approx({own: 2})


def test_circular_workbook_has_no_impact(tmp_path, monkeypatch):
    book = openpyxl.Workbook()
    ws = book.active
    ws.title = "S"
    ws["B1"] = 2
    ws["B2"] = "=B1*2.5+B3"
    ws["B3"] = "=B2/2"
    path = tmp_path / "cycle.xlsx"
    book.save(path)
    wb, graph = _load(path)
    compiled = []
    monkeypatch.setattr(recalc, "compile", lambda *args: compiled.append(args))
    constant = _input("S!B1", 2, 3, ["S!B2"])
    outputs = _outputs("S!B2")

    assert graph.cycles
    assert impact.deltas(wb, graph, constant, outputs) == []
    assert impact.combined(wb, graph, [constant], outputs) == []
    impact.fill_outputs(wb, graph, outputs)
    assert outputs[0]["value"] == 0.0
    assert compiled == []


def test_workbook_is_compiled_once(tmp_path, monkeypatch):
    wb, graph = _tiny(tmp_path)
    calls = []
    compile_ = recalc.compile
    monkeypatch.setattr(recalc, "compile", lambda *args: calls.append(args) or compile_(*args))
    outputs = _outputs("S!B2")
    units = _input("S!B1", 2, 3, ["S!B2"])
    factor = _literal(["S!B2"], 0, 2.5, 4, ["S!B2"])

    impact.fill_outputs(wb, graph, outputs)
    assert outputs[0]["value"] == 5.0
    assert _deltas((wb, graph), units, outputs) == {"S!B2": 2.5}
    assert _deltas((wb, graph), factor, outputs) == {"S!B2": 3.0}
    assert _by_output(impact.combined(wb, graph, [units, factor], outputs)) == {"S!B2": 7.0}
    assert len(calls) == 1
    assert calls[0][1] == [("S!B2", 0)]


def test_failed_compile_is_remembered(tmp_path, monkeypatch):
    wb, graph = _tiny(tmp_path)
    calls = []
    monkeypatch.setattr(recalc, "compile", lambda *args: calls.append(args))
    outputs = _outputs("S!B2")
    constant = _input("S!B1", 2, 3, ["S!B2"])

    assert impact.deltas(wb, graph, constant, outputs) == []
    assert impact.deltas(wb, graph, constant, outputs) == []
    assert impact.combined(wb, graph, [constant], outputs) == []
    impact.fill_outputs(wb, graph, outputs)
    assert outputs[0]["value"] == 0.0
    assert len(calls) == 1


def test_engine_errors_do_not_escape(tmp_path, monkeypatch):
    wb, graph = _tiny(tmp_path, {"S!B2": 5})
    outputs = _outputs("S!B2")
    constant = _input("S!B1", 2, 3, ["S!B2"])

    def broken(*args):
        raise RuntimeError("engine fell over")

    monkeypatch.setattr(recalc, "calculate", broken)
    assert impact.deltas(wb, graph, constant, outputs) == []
    assert impact.combined(wb, graph, [constant], outputs) == []
    impact.fill_outputs(wb, graph, outputs)
    assert outputs[0]["value"] == 0.0

    monkeypatch.undo()
    assert _deltas((wb, graph), constant, outputs) == {"S!B2": 2.5}
