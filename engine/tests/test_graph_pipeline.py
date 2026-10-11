"""pipeline end to end, with the checks half absent and with it stubbed, and the model date it falls back on."""

import datetime as dt
import json
import threading
from pathlib import Path

import jsonschema
import openpyxl
import pytest

import pipeline
from checks import label, news, series
from graph import parse

SCHEMA = json.loads((Path(__file__).parents[2] / "schema" / "constant.schema.json").read_text())


@pytest.fixture
def model(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Model"
    ws["A1"], ws["B1"] = "EUR/USD rate", 1.25
    ws["A2"], ws["B2"] = "Units", 100
    ws["A3"], ws["B3"] = "EBITDA", "=B2*B1*0.5"
    path = tmp_path / "model.xlsx"
    wb.save(path)
    return path


def _run(path):
    return pipeline.run(path, model_id="m1", filename="model.xlsx", built_on="2026-08-01")


def test_report_is_valid_when_the_checks_half_is_not_implemented(model, monkeypatch):
    def missing(*args):
        raise NotImplementedError

    monkeypatch.setattr(label, "apply_all", missing)
    report = _run(model)
    jsonschema.validate(report, SCHEMA)
    assert [c["cell"] for c in report["constants"]] == ["Model!B1", "Model!B2", "Model!B3"]
    assert all(c["check"] is None and c["impact"] == [] for c in report["constants"])
    # never calculated by Excel, so the value comes from the recalculation
    assert report["outputs"] == [{"cell": "Model!B3", "label": "EBITDA", "value": 62.5}]
    assert report["combined"] == []


def test_a_live_value_flows_through_to_impact(model, monkeypatch):
    def apply_all(found, wb, graph):
        for c in found:
            c["label"] = wb.context(c["cell"])["row_label"] or ""
            c["driver"] = "eurusd" if c["cell"] == "Model!B1" else None

    def check(c, built_on, as_of=None):
        return {"kind": "series", "status": "stale", "live": 1.5, "evidence": []}

    monkeypatch.setattr(label, "apply_all", apply_all)
    # stub both routes: the checks half decides whether a driver goes to series or news
    monkeypatch.setattr(series, "check", check)
    monkeypatch.setattr(news, "check", check)
    report = _run(model)
    jsonschema.validate(report, SCHEMA)

    fx = report["constants"][0]  # sorted by impact, so the only constant with one comes first
    assert fx["cell"] == "Model!B1"
    assert fx["impact"] == [{"output": "Model!B3", "delta": pytest.approx(12.5)}]
    assert report["combined"] == [{"output": "Model!B3", "delta": pytest.approx(12.5)}]
    # measured before the check ran: +1% on the rate moves EBITDA by +1%
    effect = fx["sensitivity"]["effects"][0]
    assert (effect["output"], effect["direction"], effect["pct"]) == ("Model!B3", "rises", pytest.approx(1.0))
    assert effect["path"][0]["formula"] == "=B2*B1*0.5"
    assert all("sensitivity" not in c for c in report["constants"][1:])
    assert [c["check"] for c in report["constants"][1:]] == [None, None]


def _stub_checks(monkeypatch):
    """EUR/USD is the one market constant; its check finds 1.5. Returns the constants handed to the check."""
    calls = []

    def apply_all(found, wb, graph):
        for c in found:
            c["label"] = wb.context(c["cell"])["row_label"] or ""
            c["driver"] = "eurusd" if c["cell"] == "Model!B1" else None

    def check(c, built_on, as_of=None):
        calls.append(c)
        return {"kind": "news", "status": "stale", "live": 1.5, "evidence": []}

    monkeypatch.setattr(label, "apply_all", apply_all)
    monkeypatch.setattr(news, "check", check)
    return calls


def test_start_lists_constants_before_anything_is_checked(model, monkeypatch):
    calls = _stub_checks(monkeypatch)
    report = pipeline.start(model, model_id="m1", filename="model.xlsx")
    jsonschema.validate(report, SCHEMA)
    json.dumps(report, allow_nan=False)
    assert (report["status"], report["progress"]) == ("checking", {"checked": 0, "total": 0})
    assert [c["cell"] for c in report["constants"]] == ["Model!B1", "Model!B2", "Model!B3"]
    assert all((c["label"], c["driver"], c["check"], c["impact"]) == ("", None, None, []) for c in report["constants"])
    assert report["outputs"] == [{"cell": "Model!B3", "label": "EBITDA", "value": 62.5}]
    assert report["combined"] == []
    assert report["built_on"] == report["checked_on"] == dt.date.today().isoformat()
    assert calls == []


def test_run_returns_a_finished_report(model, monkeypatch):
    calls = _stub_checks(monkeypatch)
    report = _run(model)
    jsonschema.validate(report, SCHEMA)
    json.dumps(report, allow_nan=False)
    assert (report["status"], report["progress"]) == ("done", {"checked": 1, "total": 1})
    assert report["built_on"] == "2026-08-01"
    assert report["combined"] == [{"output": "Model!B3", "delta": pytest.approx(12.5)}]
    # the check saw the sensitivity, on a copy of the constant
    assert calls[0]["sensitivity"]["effects"][0]["output"] == "Model!B3"
    assert calls[0] is not report["constants"][0]

    pipeline.finish(report)  # finished once: nothing is checked again
    assert len(calls) == 1


def test_a_model_is_finished_by_one_thread_only(model, monkeypatch):
    calls = _stub_checks(monkeypatch)
    check, entered, release = news.check, threading.Event(), threading.Event()

    def held(c, built_on, as_of=None):
        entered.set()
        assert release.wait(10)
        return check(c, built_on, as_of)

    monkeypatch.setattr(news, "check", held)
    report = pipeline.start(model, model_id="m1", filename="model.xlsx")
    lock = threading.Lock()
    first = threading.Thread(target=pipeline.finish, args=(report,), kwargs={"lock": lock})
    first.start()
    try:
        assert entered.wait(10)
        assert not lock.locked()  # not held while a check is out on the network
        pipeline.finish(report, lock=lock)  # returns at once
        assert report["status"] == "checking"
    finally:
        release.set()
        first.join(10)
    assert (report["status"], len(calls)) == ("done", 1)


def test_refresh_recalculates_one_edited_constant(model, monkeypatch):
    _stub_checks(monkeypatch)
    report = _run(model)
    fx, units = report["constants"][0], report["constants"][1]

    fx["check"] = None
    pipeline.refresh(report, fx)
    assert (fx["impact"], report["combined"]) == ([], [])

    pipeline._LOADED.clear()  # the workbook is read again from the file run was given
    fx["check"] = {"kind": "news", "status": "stale", "live": 1.0, "evidence": []}
    units["driver"], units["check"] = "units", {"kind": "news", "status": "stale", "live": 200, "evidence": []}
    pipeline.refresh(report, units)
    assert units["impact"] == [{"output": "Model!B3", "delta": pytest.approx(62.5)}]
    assert units["sensitivity"]["effects"][0]["pct"] == pytest.approx(1.0)
    assert fx["impact"] == []  # only the constant passed is recalculated on its own
    assert report["combined"] == [{"output": "Model!B3", "delta": pytest.approx(37.5)}]
    jsonschema.validate(report, SCHEMA)
    json.dumps(report, allow_nan=False)


def test_refresh_without_a_workbook_clears_the_impact(monkeypatch):
    report = {"model_id": "gone", "outputs": [], "constants": [{"cell": "Model!B1", "impact": [{"output": "x", "delta": 1}]}]}
    pipeline.refresh(report, report["constants"][0])
    assert report["constants"][0]["impact"] == []


def _dated(tmp_path, rows, sheets=()):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Cover"
    for cell, value in rows.items():
        ws[cell] = value
    for title, cells in sheets:
        extra = wb.create_sheet(title)
        for cell, value in cells.items():
            extra[cell] = value
    path = tmp_path / "dated.xlsx"
    wb.save(path)
    return parse.load(path)


@pytest.mark.parametrize(
    "row_label",
    ["Model date", "VALUATION DATE:", "Prices as of", "Balance sheet as at", "Assumptions date", "Date prepared", "Prepared on"],
)
def test_model_date_reads_a_labeled_date(tmp_path, row_label):
    wb = _dated(tmp_path, {"B2": "Fiscal year end", "C2": dt.date(2026, 12, 31), "B5": row_label, "C5": dt.datetime(2026, 8, 1)})
    assert parse.model_date(wb) == "2026-08-01"


def test_model_date_takes_the_first_in_sheet_row_column_order(tmp_path):
    wb = _dated(
        tmp_path,
        {"A9": "Valuation date", "B9": dt.date(2026, 3, 1), "A4": "As of", "D4": dt.date(2026, 2, 1), "C4": dt.date(2026, 1, 1)},
        sheets=[("DCF", {"A1": "Model date", "B1": dt.date(2025, 1, 1)})],
    )
    assert parse.model_date(wb) == "2026-01-01"


def test_model_date_is_none_without_a_labeled_date(tmp_path):
    rows = {
        "A1": "Fiscal year end", "B1": dt.date(2026, 12, 31),
        "A2": "Gas of the month", "B2": dt.date(2026, 1, 1),
        "A3": "Model date", "B3": "August 2026",
        "A4": "Valuation date", "B4": 46235,
        "B5": dt.date(2026, 5, 5),
    }
    assert parse.model_date(_dated(tmp_path, rows)) is None


def test_model_date_of_the_case_study():
    wb = parse.load(Path(__file__).parents[2] / "data" / "case_study_aug2026.xlsx")
    assert parse.model_date(wb) == "2026-08-01"


def test_built_on_comes_from_the_workbook_when_not_given(tmp_path, monkeypatch):
    _stub_checks(monkeypatch)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"], ws["B1"] = "As of", dt.date(2025, 6, 30)
    ws["A2"], ws["B2"] = "Units", 100
    ws["A3"], ws["B3"] = "EBITDA", "=B2*2"
    path = tmp_path / "dated.xlsx"
    wb.save(path)
    assert pipeline.run(path, model_id="m2", filename="dated.xlsx")["built_on"] == "2025-06-30"
    assert pipeline.run(path, model_id="m2", filename="dated.xlsx", built_on="2024-01-01")["built_on"] == "2024-01-01"
