"""impact.sensitivity: which way and how far each output moves when a constant rises a little."""

import json
from pathlib import Path

import openpyxl
import pytest

from graph import constants, impact, parse
from graph.parse import Graph

DATA = Path(__file__).parents[2] / "data"


@pytest.fixture(scope="module")
def model(tmp_path_factory):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Model"
    ws["A1"], ws["B1"] = "Price", 10
    ws["A2"], ws["B2"] = "FX rate", 2.0
    ws["A3"], ws["B3"] = "Tariff", 0
    ws["B3"].number_format = "0.0%"
    ws["A4"], ws["B4"] = "Cost", "=100/B2*(1+B3)"
    ws["A5"], ws["B5"], ws["C5"] = "Overhead", "=B1*0.5", "=B1*0.5"
    ws["A6"], ws["B6"] = "EBITDA", "=B1*5-B4-B5-C5"
    ws["A7"], ws["B7"] = "Unused", 99
    path = tmp_path_factory.mktemp("sensitivity") / "model.xlsx"
    wb.save(path)
    loaded = parse.load(path)
    graph = parse.build_graph(loaded)
    return loaded, graph, {c["cell"]: c for c in constants.find(loaded, graph)}


def _effect(result, label="EBITDA"):
    return next(e for e in result["effects"] if e["label"] == label)


def test_direct_driver_raises_the_output(model):
    wb, graph, found = model
    result = impact.sensitivity(wb, graph, found["Model!B1"])
    assert result["bump"] == {"from": 10, "to": 10.1, "kind": "relative", "size": 0.01}
    effect = _effect(result)
    # 10 -> 10.1: revenue +0.5, two overhead cells -0.05 each; baseline 50 - 50 - 5 - 5 = -10
    assert effect["direction"] == "rises"
    assert effect["delta"] == pytest.approx(0.4)
    assert effect["baseline"] == pytest.approx(-10)
    assert effect["pct"] == pytest.approx(4.0)


def test_divisor_moves_the_output_the_other_way_and_the_path_shows_why(model):
    wb, graph, found = model
    effect = _effect(impact.sensitivity(wb, graph, found["Model!B2"]))
    assert effect["direction"] == "rises"  # a higher rate divides the cost down
    assert effect["delta"] == pytest.approx(50 - 100 / 2.02, abs=1e-5)
    cost, ebitda = effect["path"]
    assert (cost["cell"], cost["label"], cost["formula"], cost["direction"]) == ("Model!B4", "Cost", "=100/B2*(1+B3)", "falls")
    assert (ebitda["cell"], ebitda["direction"]) == ("Model!B6", "rises")


def test_zero_percent_constant_is_bumped_by_one_point(model):
    wb, graph, found = model
    result = impact.sensitivity(wb, graph, found["Model!B3"])
    assert result["bump"] == {"from": 0, "to": 0.01, "kind": "absolute", "size": 0.01}
    effect = _effect(result)
    assert effect["direction"] == "falls"
    assert effect["delta"] == pytest.approx(-0.5)


def test_embedded_group_is_bumped_together(model):
    wb, graph, found = model
    effect = _effect(impact.sensitivity(wb, graph, found["Model!B5:C5"]))
    assert effect["delta"] == pytest.approx(-2 * 10 * 0.005)  # both overhead cells move
    assert effect["path"][0]["formula"] == "=B1*0.5"


def test_outputs_argument_is_optional_and_the_result_is_json(model):
    wb, graph, found = model
    outputs = constants.find_outputs(wb, graph)
    assert impact.sensitivity(wb, graph, found["Model!B1"], outputs) == impact.sensitivity(wb, graph, found["Model!B1"])
    json.dumps(impact.sensitivity(wb, graph, found["Model!B1"]), allow_nan=False)


def test_constant_that_reaches_no_output_has_no_sensitivity(model):
    wb, graph, _ = model
    assert impact.sensitivity(wb, graph, {"cell": "Model!B7", "value": 99, "embedded": False, "outputs": []}) is None


def test_without_a_recalculation_the_path_and_formulas_are_still_given(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Model"
    ws["A1"], ws["B1"] = "Rate", 0.05
    ws["A2"], ws["B2"] = "Cost", "=100*B1+C2"
    ws["C2"] = "=B2*0"  # circular, so the workbook is not recalculated
    ws["A3"], ws["B3"] = "EBITDA", "=1000-B2"
    path = tmp_path / "cycle.xlsx"
    wb.save(path)
    loaded = parse.load(path)
    graph = parse.build_graph(loaded)
    constant = next(c for c in constants.find(loaded, graph) if c["cell"] == "Model!B1")
    result = impact.sensitivity(loaded, graph, constant)
    assert result["bump"] is None
    (effect,) = result["effects"]
    assert (effect["delta"], effect["direction"]) == (None, None)
    assert [(s["cell"], s["formula"]) for s in effect["path"]] == [("Model!B2", "=100*B1+C2"), ("Model!B3", "=1000-B2")]


def test_routes():
    graph = Graph(dependents={"A": {"B", "C"}, "B": {"D"}, "C": {"D"}, "D": {"E"}, "X": {"E"}})
    assert graph.routes(["A"], {"E", "Z"}) == {"E": ["A", "B", "D", "E"]}
    assert graph.routes(["A", "X"], {"E"}) == {"E": ["X", "E"]}
    assert graph.routes(["D"], {"D"}) == {"D": ["D"]}


CASES = [
    # (workbook, constant, value, output label, direction when the constant rises)
    ("case_study_aug2026.xlsx", "Assumptions!B20", 96.95, "EBITDA", "falls"),  # oil is a cost
    ("case_study_aug2026.xlsx", "Assumptions!B21", 1.1519, "EBITDA", "falls"),  # a stronger euro costs more dollars
    ("case_study_airline.xlsx", "P&L!C14:G14", 142.63, "EBITDA", "falls"),  # yen per dollar divides yen revenue
    ("case_study_trucking.xlsx", "P&L!C27:G27", 17.2207, "EBITDA", "rises"),  # pesos per dollar divides peso costs
    ("case_study_coffee.xlsx", "P&L!C16:G16", 4.9792, "EBITDA", "rises"),  # reais per dollar divides real costs
    ("case_study_airline.xlsx", "Assumptions!B33", 0.0441, "Net income", "falls"),  # SOFR raises interest
]


@pytest.mark.parametrize("workbook, cell, value, label, direction", CASES)
def test_direction_on_the_case_studies(workbook, cell, value, label, direction):
    if not (DATA / workbook).exists():
        pytest.skip("case study workbook not present")
    wb = parse.load(DATA / workbook)
    graph = parse.build_graph(wb)
    constant = next(c for c in constants.find(wb, graph) if c["cell"] == cell and c["value"] == pytest.approx(value))
    effect = _effect(impact.sensitivity(wb, graph, constant), label)
    assert effect["direction"] == direction
    assert effect["path"] and all(step["formula"].startswith("=") for step in effect["path"])
