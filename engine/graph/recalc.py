"""Recalculate a workbook with some cells overridden. The only module that imports `formulas`:
swapping the calculation engine means rewriting this file."""

import math
import numbers
import re
import tempfile
from dataclasses import dataclass, field
from pathlib import Path

import formulas
import openpyxl
from openpyxl.formula.tokenizer import Token, Tokenizer
from openpyxl.worksheet.formula import ArrayFormula, DataTableFormula

from graph.refs import make_key, split_key

# compiling costs about half a second per 1,000 formulas and 10 seconds per million cells of used area
MAX_FORMULAS = 5_000
MAX_AREA = 500_000

_FILENAME = "model.xlsx"
_SCRATCH = "_drift"
_UNSUPPORTED_IN_SHEET_NAME = re.compile(r"['%]")  # `formulas` cannot follow references into such a sheet


@dataclass
class Model:
    """A compiled workbook."""

    engine: object
    scratch: dict[tuple[str, int], str]  # (formula cell, literal index) -> key of the cell now holding that number
    aliases: dict[str, str] = field(default_factory=dict)  # sheet name -> its name in the compiled copy


def too_big(wb) -> bool:
    """True when compiling the workbook would take too long."""
    area = sum((bottom - top + 1) * (right - left + 1) for top, left, bottom, right in wb.bounds.values())
    return len(wb.formulas) > MAX_FORMULAS or area > MAX_AREA


def compile(wb, literals) -> Model | None:
    """Model of the workbook with the given (formula cell, literal index) numbers hoisted into cells
    of their own, so they can be overridden. None when the workbook is too big or anything fails."""
    if too_big(wb):
        return None
    try:
        with tempfile.TemporaryDirectory(prefix="drift-") as directory:
            path = Path(directory) / _FILENAME
            scratch = hoist(wb, literals, path)
            engine = formulas.ExcelModel().loads(str(path)).finish()
        return Model(engine, scratch, _names(wb.sheets)[0])
    except Exception:
        return None


def calculate(model: Model, overrides: dict[str, float], cells: list[str]) -> dict[str, float | None]:
    """Value of each cell once the overridden cells hold their new values. None where the result is
    not a finite number, and everywhere when an override names a cell the model does not have."""
    result: dict[str, float | None] = dict.fromkeys(cells)
    try:
        inputs = {_node(model, cell): value for cell, value in overrides.items()}
        if not cells or any(node not in model.engine.cells for node in inputs):
            return result
        # asking the engine for specific outputs can be several times slower than a full recalculation
        solution = model.engine.calculate(inputs=inputs)
        for cell in cells:
            node = _node(model, cell)
            if node in solution:
                result[cell] = _number(solution[node].value[0, 0])
    except Exception:
        return dict.fromkeys(cells)
    return result


def hoist(wb, literals, dest) -> dict[tuple[str, int], str]:
    """Save a copy of the workbook to dest in which each (formula cell, literal index) number sits in
    its own cell on a scratch sheet, a leading - and trailing % folded in, and the formula points at it:
    =B2*(1+5%) becomes =B2*(1+_drift!$A$1) with 0.05 in _drift!A1. Returns literal -> scratch cell key.
    Numbers in array constants and array formulas cannot be replaced: they stay, and are not returned."""
    book = openpyxl.load_workbook(wb.path)
    try:
        aliases, scratch_name = _names(book.sheetnames)
        scratch = book.create_sheet(scratch_name)
        wanted: dict[str, set[int]] = {}
        for key, index in literals:
            wanted.setdefault(key, set()).add(index)

        hoisted: dict[tuple[str, int], str] = {}
        for key, indices in wanted.items():
            sheet, coord = split_key(key)
            cell = book[sheet][coord] if sheet in book.sheetnames else None
            if cell is None or cell.data_type != "f" or not isinstance(cell.value, str):
                continue
            cell.value, values = _rewrite(cell.value, indices, scratch_name, len(hoisted) + 1)
            for index, value in values.items():
                row = len(hoisted) + 1
                scratch.cell(row=row, column=1, value=value)
                hoisted[key, index] = make_key(scratch_name, f"A{row}")

        _drop_data_tables(book)
        if aliases:
            _rename_sheets(book, aliases)
        book.save(dest)
        return hoisted
    finally:
        book.close()


def _rewrite(formula: str, indices: set[int], scratch: str, first_row: int) -> tuple[str, dict[int, float]]:
    """The formula with its index-th number literals replaced by scratch!$A$first_row, the next row, ...
    and the value each replaced literal had, in that order."""
    tokens = Tokenizer(formula).items
    values: dict[int, float] = {}
    dropped: set[int] = set()
    array_depth = 0
    index = -1
    for position, token in enumerate(tokens):
        if token.type == Token.ARRAY:
            array_depth += 1 if token.subtype == Token.OPEN else -1
        if token.type != Token.OPERAND or token.subtype != Token.NUMBER:
            continue
        index += 1
        value = float(token.value)
        if index not in indices or array_depth or not math.isfinite(value):
            continue
        before = _neighbour(tokens, position, -1)
        if before is not None and tokens[before].type == Token.OP_PRE and tokens[before].value == "-":
            value = -value
            dropped.update(range(before, position))
        after = _neighbour(tokens, position, 1)
        if after is not None and tokens[after].type == Token.OP_POST and tokens[after].value == "%":
            value /= 100
            dropped.update(range(position + 1, after + 1))
        token.value = f"{scratch}!$A${first_row + len(values)}"
        values[index] = value
    return "=" + "".join(token.value for position, token in enumerate(tokens) if position not in dropped), values


def _neighbour(tokens: list[Token], position: int, step: int) -> int | None:
    """Position of the nearest token before (step -1) or after (step 1) that is not whitespace."""
    position += step
    while 0 <= position < len(tokens):
        if tokens[position].type != Token.WSPACE:
            return position
        position += step
    return None


def _names(sheets: list[str]) -> tuple[dict[str, str], str]:
    """Replacement names for the sheets `formulas` cannot reference, and a free name for the scratch sheet."""
    taken = {sheet.lower() for sheet in sheets}

    def free(name: str, suffix: str) -> str:
        while name.lower() in taken:
            name += suffix
        taken.add(name.lower())
        return name

    aliases = {
        sheet: free(_UNSUPPORTED_IN_SHEET_NAME.sub("_", sheet), "_")
        for sheet in sheets
        if _UNSUPPORTED_IN_SHEET_NAME.search(sheet)
    }
    return aliases, free(_SCRATCH, "_")


def _drop_data_tables(book) -> None:
    """Empty the cells of what-if data tables, which `formulas` cannot load. parse.load skips them too."""
    for ws in book.worksheets:
        for cell in ws._cells.values():
            if isinstance(cell.value, DataTableFormula):
                cell.value = None


def _rename_sheets(book, aliases: dict[str, str]) -> None:
    """Rename sheets and repoint every formula and defined name that mentions them."""
    quoted = {"'" + old.replace("'", "''").lower() + "'!": f"'{new}'!" for old, new in aliases.items()}
    pattern = re.compile("|".join(map(re.escape, quoted)), re.IGNORECASE)

    def repoint(text: str) -> str:
        return pattern.sub(lambda match: quoted[match[0].lower()], text)

    for ws in book.worksheets:
        for cell in ws._cells.values():
            if isinstance(cell.value, ArrayFormula):
                if isinstance(cell.value.text, str):
                    cell.value.text = repoint(cell.value.text)
            elif cell.data_type == "f" and isinstance(cell.value, str):
                cell.value = repoint(cell.value)
        for defined in ws.defined_names.values():
            if isinstance(defined.attr_text, str):
                defined.attr_text = repoint(defined.attr_text)
    for defined in book.defined_names.values():
        if isinstance(defined.attr_text, str):
            defined.attr_text = repoint(defined.attr_text)
    for old, new in aliases.items():
        book[old].title = new


def _node(model: Model, cell: str) -> str:
    """"P&L!C27" -> "'[model.xlsx]P&L'!C27", the engine's name for the cell."""
    sheet, coord = split_key(cell)
    return f"'[{_FILENAME}]{model.aliases.get(sheet, sheet).upper()}'!{coord}"


def _number(value: object) -> float | None:
    """A finite number as a float. None for errors, text, booleans and empty cells."""
    if isinstance(value, bool) or not isinstance(value, numbers.Real):
        return None
    value = float(value)
    return value if math.isfinite(value) else None
