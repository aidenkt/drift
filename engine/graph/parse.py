"""Load a workbook and build the cell dependency graph."""

from collections import deque
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl
from openpyxl.worksheet.formula import ArrayFormula, DataTableFormula

from graph.refs import Bounds, make_key


@dataclass
class Workbook:
    """Everything read from the file, keyed by "Sheet!A1"."""

    path: Path
    sheets: list[str]
    formulas: dict[str, str] = field(default_factory=dict)  # formula cells only, e.g. "=B5*1.08"
    values: dict[str, object] = field(default_factory=dict)  # literal, or Excel's cached result (None if never calculated)
    formats: dict[str, str] = field(default_factory=dict)  # Excel number format string
    comments: dict[str, str] = field(default_factory=dict)
    bounds: dict[str, Bounds] = field(default_factory=dict)  # used area per sheet
    names: dict[str, str] = field(default_factory=dict)  # defined name (upper case) -> reference text

    def context(self, cell: str) -> dict:
        from graph.context import context_for

        return context_for(self, cell)


@dataclass
class Literal:
    """A number written inside a formula."""

    cell: str  # the formula cell
    index: int  # position among the number literals in that formula
    value: float


@dataclass
class Graph:
    precedents: dict[str, set[str]] = field(default_factory=dict)  # cell -> cells it reads
    dependents: dict[str, set[str]] = field(default_factory=dict)  # cell -> cells that read it
    literals: dict[str, list[Literal]] = field(default_factory=dict)
    cycles: list[list[str]] = field(default_factory=list)

    def downstream(self, cell: str) -> set[str]:
        """Every cell that depends on this one, transitively. Excludes the cell itself."""
        seen: set[str] = set()
        queue = deque([cell])
        while queue:
            for dependent in self.dependents.get(queue.popleft(), ()):
                if dependent not in seen:
                    seen.add(dependent)
                    queue.append(dependent)
        seen.discard(cell)
        return seen


def load(path) -> Workbook:
    """openpyxl twice: formulas (data_only=False) and cached values (data_only=True)."""
    path = Path(path)
    formula_wb = openpyxl.load_workbook(path, data_only=False)
    value_wb = openpyxl.load_workbook(path, data_only=True)
    try:
        wb = Workbook(path=path, sheets=[ws.title for ws in formula_wb.worksheets])
        for ws, value_ws in zip(formula_wb.worksheets, value_wb.worksheets):
            _read_sheet(wb, ws, value_ws)
        for name, defined in formula_wb.defined_names.items():
            if defined.localSheetId is None and not name.startswith("_xlnm."):
                wb.names[name.upper()] = defined.attr_text
        return wb
    finally:
        formula_wb.close()
        value_wb.close()


def _read_sheet(wb: Workbook, ws, value_ws) -> None:
    """Copy one worksheet into wb; value_ws is the same sheet from the data_only workbook."""
    sheet = ws.title
    wb.bounds[sheet] = (ws.min_row, ws.min_column, ws.max_row, ws.max_column)
    cached = value_ws._cells
    # _cells holds only the cells stored in the file; iter_rows() would create every blank in the used rectangle
    for position, cell in ws._cells.items():
        value = cell.value
        if isinstance(value, DataTableFormula):
            continue
        comment = cell.comment
        if value is None and comment is None:
            continue
        key = make_key(sheet, cell.coordinate)
        if comment is not None:
            wb.comments[key] = comment.text
        if value is None:
            continue
        if cell.data_type == "f":
            wb.formulas[key] = _formula_text(value)
            cached_cell = cached.get(position)
            wb.values[key] = cached_cell.value if cached_cell is not None else None
        else:
            wb.values[key] = value
        wb.formats[key] = cell.number_format


def _formula_text(value) -> str:
    text = value.text if isinstance(value, ArrayFormula) else value
    return text if text.startswith("=") else f"={text}"


def build_graph(wb):
    """Cell -> cells it references, from formula refs (A1, ranges, cross-sheet)."""
    raise NotImplementedError
