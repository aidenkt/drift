"""What sits around a cell: the raw text a labeler needs. No interpretation."""

from datetime import date, datetime

from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.utils.cell import coordinate_from_string

from graph.refs import make_key, split_key


def context_for(wb, cell: str) -> dict:
    """{"sheet", "row_label", "col_header", "number_format", "comment"} for a cell key."""
    sheet, coord = split_key(cell)
    if sheet not in wb.bounds:
        raise KeyError(sheet)
    key = make_key(sheet, coord)
    letter, row = coordinate_from_string(split_key(key)[1])
    col = column_index_from_string(letter)
    comment = (wb.comments.get(key) or "").strip()
    return {
        "sheet": sheet,
        "row_label": _row_label(wb, sheet, row, col),
        "col_header": _col_header(wb, sheet, row, col),
        "number_format": wb.formats.get(key, "General"),
        "comment": comment or None,
    }


def _row_label(wb, sheet: str, row: int, col: int) -> str | None:
    """Nearest text to the left in the row, skipping numbers, dates and blanks."""
    _, min_col, _, max_col = wb.bounds[sheet]
    for c in range(min(col - 1, max_col), min_col - 1, -1):
        text = _text(wb.values.get(make_key(sheet, f"{get_column_letter(c)}{row}")))
        if text:
            return text
    return None


def _col_header(wb, sheet: str, row: int, col: int) -> str | None:
    """Nearest text above in the column, else a date or year at the top of the column."""
    min_row, _, max_row, _ = wb.bounds[sheet]
    letter = get_column_letter(col)
    top = None
    for r in range(min(row - 1, max_row), min_row - 1, -1):
        key = make_key(sheet, f"{letter}{r}")
        value = wb.values.get(key)
        text = _text(value)
        if text:
            return text
        if isinstance(value, str):
            continue
        if value is not None:
            top = value
        elif key in wb.formulas:
            top = None  # a never-calculated formula is the top cell, but says nothing
    return _period(top)


def _text(value: object) -> str | None:
    if isinstance(value, str):
        return value.strip() or None
    return None


def _period(value: object) -> str | None:
    """A date as YYYY-MM-DD, a plausible year as "2026", anything else None."""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, bool):
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    if isinstance(value, int) and 1900 <= value <= 2100:
        return str(value)
    return None
